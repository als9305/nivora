import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles & Colors
    primary_color = RGBColor(14, 116, 144)   # Cyan 700
    dark_color = RGBColor(15, 23, 42)       # Slate 900
    accent_color = RGBColor(2, 132, 199)     # Sky 600

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("🛡️ DEEPGUARD — HACKATHON MASTER DOCUMENTATION")
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = primary_color

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Deepfake-Resilient Adaptive Multi-Modal Authentication Framework\nSmart India Hackathon (SIH) — Cybersecurity & Banking Track")
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = dark_color

    doc.add_paragraph()

    # Section 1: Executive Summary
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.runs[0].font.color.rgb = primary_color

    p_exec = doc.add_paragraph(
        "DeepGuard is a next-generation cybersecurity authentication platform designed to protect banks, "
        "financial institutions, and digital identity systems against AI-generated deepfakes, photo replay attacks, and synthetic voice clones.\n\n"
        "Traditional banking apps rely on password + single-factor face recognition. Today, attackers can easily bypass "
        "face recognition using cheap AI deepfake apps or printed photos. DeepGuard solves this by combining 5 independent security "
        "signals into a dynamic 0–100 risk engine that automatically grants access, requests Step-Up OTP, or blocks security threats in real-time."
    )
    p_exec.style.font.size = Pt(10.5)

    # Section 2: Why This Is Needed
    h2 = doc.add_heading("2. Why This Project is Needed (Problem Statement)", level=1)
    h2.runs[0].font.color.rgb = primary_color

    items = [
        ("The Rise of Generative AI Deepfakes: ", "Tools like DeepFaceLab, ElevenLabs, and FaceSwap allow fraudsters to create realistic video and voice clones of victims in minutes."),
        ("Banking & Financial Fraud: ", "Cybercriminals use deepfakes during online Video KYC customer onboarding, net banking logins, and high-value wire transfers to steal funds."),
        ("Flaw of Single-Factor Biometrics: ", "Traditional face match checks identity, but does NOT check whether the video feed is a real living human or a computer-generated synthetic video."),
        ("The Convenience vs Security Balance: ", "Banks need high security, but users demand convenience. DeepGuard provides frictionless 1-click login for genuine users, and only challenges suspicious logins.")
    ]
    for b_title, b_desc in items:
        p = doc.add_paragraph(style='List Bullet')
        r1 = p.add_run(b_title)
        r1.bold = True
        r1.font.color.rgb = dark_color
        p.add_run(b_desc)

    # Section 3: What We Have Built
    h3 = doc.add_heading("3. What We Have Built (Solution Overview)", level=1)
    h3.runs[0].font.color.rgb = primary_color

    doc.add_paragraph("We have engineered a complete, working full-stack software prototype consisting of 8 core modules:")

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    hdr_cells = table.rows[0].cells
    headers = ["Module", "What It Does", "Technical Mechanism"]
    widths = [Inches(1.8), Inches(2.3), Inches(2.7)]

    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "0E7490")
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(10)

    modules_data = [
        ("1. Face Identity Match", "Compares user face against baseline.", "OpenCV 128-d spatial gradient vector comparison."),
        ("2. Active Liveness", "Verifies physical motion (turn, smile, blink).", "Landmark vector variance. Fails static photo prints."),
        ("3. Deepfake Inspection", "Detects synthetic AI video face swaps.", "Multi-frame FFT spectral residual & color gamut analysis."),
        ("4. Voice Anti-Spoofing", "Verifies speaker similarity & blocks TTS clones.", "16-band energy spectrum extraction & pitch stability."),
        ("5. Device & Session Risk", "Fingerprints browser session & automation tools.", "User-Agent inspection & session risk scoring."),
        ("6. Composite Risk Engine", "Calculates unified 0–100 threat score.", "Weighted multi-modal risk algorithm (ALLOW, STEP-UP, BLOCK)."),
        ("7. AI Explainability Agent", "Generates human-readable security reports.", "Audit evidence reports explaining why login was allowed or blocked."),
        ("8. Attack Simulator Studio", "Live demonstration studio for judges.", "Simulates 6 attack scenarios (Genuine, Photo, Deepfake, Voice, Replay, Borderline).")
    ]

    for row_idx, (m_name, m_does, m_tech) in enumerate(modules_data):
        row_cells = table.add_row().cells
        bg_hex = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for i, val in enumerate([m_name, m_does, m_tech]):
            row_cells[i].text = val
            row_cells[i].width = widths[i]
            set_cell_background(row_cells[i], bg_hex)
            for p in row_cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)

    doc.add_paragraph()

    # Section 4: Tech Stack & System Architecture
    h4 = doc.add_heading("4. How It Is Built (Tech Stack & Architecture)", level=1)
    h4.runs[0].font.color.rgb = primary_color

    arch_items = [
        ("Frontend: ", "React, Vite, Vanilla CSS, Tailwind CSS, Recharts Telemetry Charts, WebRTC Media Stream."),
        ("Backend: ", "FastAPI (Python 3.12), OpenCV ML Module, NumPy FFT Spectral Engine, Passlib/Bcrypt, PyJWT."),
        ("Database: ", "Primary local MongoDB (deepguard_db.users & auth_logs) with a Zero-Downtime Thread-Safe In-Memory DB Fallback."),
        ("Dual Input System: ", "Supports both Live WebRTC Camera/Mic Streams and Drag-and-Drop / Click-to-Attach File Uploads (Images: .png/.jpg, Videos: .mp4/.webm, Audio: .wav/.mp3).")
    ]
    for a_title, a_desc in arch_items:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(a_title)
        r.bold = True
        p.add_run(a_desc)

    # Section 5: How the System Works
    h5 = doc.add_heading("5. How The System Works (Step-by-Step Flow)", level=1)
    h5.runs[0].font.color.rgb = primary_color

    doc.add_paragraph("Phase 1: Registration (Onboarding)").runs[0].bold = True
    doc.add_paragraph(
        "User registers with Name, Email, and Password. Camera captures face -> converts to non-reversible 128-d vector. "
        "Microphone records voice phrase -> extracts 16-band energy spectrum vector. "
        "Privacy Protection: Raw images and audio are discarded immediately; only mathematical vectors are stored in MongoDB (GDPR / DPDP Compliant)."
    )

    doc.add_paragraph("Phase 2: Live Adaptive Authentication (Login Pipeline)").runs[0].bold = True
    doc.add_paragraph("Step 0: Password check -> Issues randomized motion challenge (e.g. TURN_LEFT) and phrase.")
    doc.add_paragraph("Step 1: Face Vector Comparison (20% Weight)")
    doc.add_paragraph("Step 2: Active Liveness Motion Evaluation (25% Weight)")
    doc.add_paragraph("Step 3: Deepfake Temporal Frequency Inspection (30% Weight)")
    doc.add_paragraph("Step 4: Voice Verification & Anti-Spoofing (15% Weight)")
    
    p_math = doc.add_paragraph()
    r_math = p_math.add_run("Composite Risk Score = 0.20(Face) + 0.25(Liveness) + 0.30(Deepfake) + 0.15(Voice) + 0.10(Device)")
    r_math.bold = True
    r_math.font.color.rgb = primary_color

    doc.add_paragraph("Decision Boundaries:")
    doc.add_paragraph("• LOW RISK (0 – 35) -> ALLOW: Instant 1-click login.")
    doc.add_paragraph("• MEDIUM RISK (36 – 70) -> STEP-UP OTP: Prompts user for a 6-digit authorization code.")
    doc.add_paragraph("• HIGH RISK (71 – 100) -> BLOCK: Access blocked immediately & logged in Security Dashboard.")

    # Section 6: SIH Differentiators
    h6 = doc.add_heading("6. Key Differentiators for SIH Hackathon", level=1)
    h6.runs[0].font.color.rgb = primary_color

    sih_diffs = [
        ("100% Problem Statement Coverage: ", "Includes liveness, deepfake detection, multimodal scoring, risk-based authentication, explainability, and privacy compliance."),
        ("Zero-Downtime Resilience: ", "If MongoDB is stopped or a hardware camera is unplugged, DeepGuard automatically uses built-in fallbacks. It will never crash during a judge presentation."),
        ("Attack Simulator Studio (/simulation): ", "Judges can test deepfakes and voice spoofs with 1 click without needing hardware setup."),
        ("Privacy-Aware Architecture: ", "Demonstrates real compliance with India's DPDP Act 2023 by storing vector embeddings instead of raw biometric media.")
    ]
    for d_title, d_desc in sih_diffs:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(d_title)
        r.bold = True
        p.add_run(d_desc)

    # Section 7: Quick Start Guide
    h7 = doc.add_heading("7. Quick Start Guide for Team Members", level=1)
    h7.runs[0].font.color.rgb = primary_color

    doc.add_paragraph("Step 1: Start Backend Server").runs[0].bold = True
    doc.add_paragraph("cd /d D:\\projects\\DeepGuard\\backend\nset PYTHONPATH=.\n.\\.venv-ml\\Scripts\\python.exe -m uvicorn app.main:app --port 8000")

    doc.add_paragraph("Step 2: Start Frontend App").runs[0].bold = True
    doc.add_paragraph("cd /d D:\\projects\\DeepGuard\\frontend\nnpm run dev")

    doc.add_paragraph("Step 3: Open URLs").runs[0].bold = True
    doc.add_paragraph("• Main Web App: http://localhost:5173\n• Attack Simulator Studio: http://localhost:5173/simulation\n• Security Dashboard: http://localhost:5173/dashboard")

    out_path = "D:\\projects\\DeepGuard\\DeepGuard_SIH_Master_Documentation.docx"
    doc.save(out_path)
    print(f"Document created successfully at: {out_path}")

if __name__ == "__main__":
    create_document()
